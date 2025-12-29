from .base_parser import BaseSOPParser
import io


class DocxSOPParser(BaseSOPParser):
    def extract_text(self, source) -> str:
        from docx import Document
        
        if isinstance(source, bytes):
            doc = Document(io.BytesIO(source))
        else:
            doc = Document(source)
        
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        
        return '\n'.join(paragraphs)
